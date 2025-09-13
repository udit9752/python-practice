from __future__ import annotations

import io
from docx import Document
from docx.shared import Pt
from docx.enum.style import WD_STYLE_TYPE


def render_text_to_docx(optimized_text: str) -> bytes:
	"""Render plain-text optimized resume into a minimal DOCX.
	Heuristics:
	- Lines in ALL CAPS and without bullets are treated as headings
	- Bullet lines starting with "- " are kept as list items (simple paragraphs with hyphen)
	"""
	doc = Document()
	_styles = doc.styles
	if "Heading ATS" not in _styles:
		heading_style = _styles.add_style("Heading ATS", WD_STYLE_TYPE.PARAGRAPH)
		heading_style.font.size = Pt(12)
		heading_style.font.bold = True
		# Use a simple font
		heading_style.font.name = "Calibri"
	if "Body ATS" not in _styles:
		body_style = _styles.add_style("Body ATS", WD_STYLE_TYPE.PARAGRAPH)
		body_style.font.size = Pt(11)
		body_style.font.name = "Calibri"

	for raw_line in optimized_text.split("\n"):
		line = raw_line.strip()
		if not line:
			doc.add_paragraph("", style="Body ATS")
			continue
		is_bullet = line.startswith("- ")
		is_heading = (not is_bullet) and line.isupper() and len(line) <= 48
		if is_heading:
			p = doc.add_paragraph(line, style="Heading ATS")
		else:
			p = doc.add_paragraph(line, style="Body ATS")

	bio = io.BytesIO()
	doc.save(bio)
	return bio.getvalue()