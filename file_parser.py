import io
import logging
from typing import Tuple
import PyPDF2
from docx import Document
from config import Config

logger = logging.getLogger(__name__)

class FileParser:
    """Handles parsing of different file formats for resume content extraction."""
    
    def __init__(self):
        self.supported_formats = Config.SUPPORTED_FORMATS
        self.max_file_size = Config.MAX_FILE_SIZE
    
    def parse_file(self, file_content: bytes, filename: str) -> Tuple[str, str]:
        """
        Parse file content and extract text.
        
        Args:
            file_content: Raw file content as bytes
            filename: Name of the file to determine format
            
        Returns:
            Tuple of (extracted_text, file_format)
            
        Raises:
            ValueError: If file format is not supported or file is too large
            Exception: If parsing fails
        """
        # Check file size
        if len(file_content) > self.max_file_size:
            raise ValueError(f"File size exceeds maximum allowed size of {self.max_file_size} bytes")
        
        # Determine file format
        file_extension = filename.lower().split('.')[-1] if '.' in filename else ''
        
        if file_extension not in self.supported_formats:
            raise ValueError(f"Unsupported file format: {file_extension}. Supported formats: {self.supported_formats}")
        
        try:
            if file_extension == 'pdf':
                return self._parse_pdf(file_content), 'pdf'
            elif file_extension == 'docx':
                return self._parse_docx(file_content), 'docx'
            elif file_extension == 'txt':
                return self._parse_txt(file_content), 'txt'
        except Exception as e:
            logger.error(f"Error parsing {file_extension} file: {str(e)}")
            raise Exception(f"Failed to parse {file_extension} file: {str(e)}")
    
    def _parse_pdf(self, file_content: bytes) -> str:
        """Extract text from PDF file."""
        try:
            pdf_file = io.BytesIO(file_content)
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            
            text = ""
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
            
            if not text.strip():
                raise Exception("No text could be extracted from the PDF")
            
            return text.strip()
        except Exception as e:
            raise Exception(f"PDF parsing error: {str(e)}")
    
    def _parse_docx(self, file_content: bytes) -> str:
        """Extract text from DOCX file."""
        try:
            docx_file = io.BytesIO(file_content)
            doc = Document(docx_file)
            
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
            
            if not text.strip():
                raise Exception("No text could be extracted from the DOCX file")
            
            return text.strip()
        except Exception as e:
            raise Exception(f"DOCX parsing error: {str(e)}")
    
    def _parse_txt(self, file_content: bytes) -> str:
        """Extract text from TXT file."""
        try:
            # Try different encodings
            encodings = ['utf-8', 'utf-16', 'latin-1', 'cp1252']
            
            for encoding in encodings:
                try:
                    text = file_content.decode(encoding)
                    if text.strip():
                        return text.strip()
                except UnicodeDecodeError:
                    continue
            
            raise Exception("Could not decode text file with any supported encoding")
        except Exception as e:
            raise Exception(f"TXT parsing error: {str(e)}")
    
    def validate_content(self, text: str) -> bool:
        """
        Validate that extracted text appears to be resume content.
        
        Args:
            text: Extracted text content
            
        Returns:
            True if content appears valid, False otherwise
        """
        if not text or len(text.strip()) < 50:
            return False
        
        # Basic heuristics to check if it looks like a resume
        resume_keywords = [
            'experience', 'education', 'skills', 'work', 'employment',
            'university', 'college', 'degree', 'certificate', 'project',
            'email', 'phone', 'address', 'objective', 'summary'
        ]
        
        text_lower = text.lower()
        keyword_count = sum(1 for keyword in resume_keywords if keyword in text_lower)
        
        return keyword_count >= 2  # At least 2 resume-related keywords